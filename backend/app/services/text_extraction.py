from pathlib import Path

import docx
import openpyxl
from pypdf import PdfReader

# Shared by both course-material and student-submission uploads. A statistics
# course's real work isn't just PDFs/text — spreadsheets (regression output,
# pivot tables) and Word docs are common submission formats. Maps each
# accepted content type to the file extension used when storing the upload.
SUPPORTED_CONTENT_TYPES: dict[str, str] = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}

# Browsers are inconsistent about the Content-Type they report for CSV/Excel
# uploads (Windows Chrome commonly sends "application/vnd.ms-excel" for a
# plain .csv; some send "application/octet-stream" for anything unfamiliar).
# Fall back to the filename's extension rather than rejecting a real file
# over a MIME-sniffing quirk.
_EXTENSION_TO_CONTENT_TYPE: dict[str, str] = {ext: ct for ct, ext in SUPPORTED_CONTENT_TYPES.items()}


def resolve_upload_type(content_type: str | None, filename: str | None) -> tuple[str, str] | None:
    """Returns (normalized_content_type, extension) for a supported upload,
    or None if neither the declared content type nor the filename extension
    is recognized."""
    if content_type in SUPPORTED_CONTENT_TYPES:
        return content_type, SUPPORTED_CONTENT_TYPES[content_type]
    suffix = Path(filename).suffix.lower() if filename else ""
    if suffix in _EXTENSION_TO_CONTENT_TYPE:
        return _EXTENSION_TO_CONTENT_TYPE[suffix], suffix
    return None


def _extract_xlsx_text(file_path: str) -> str:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_docx_text(file_path: str) -> str:
    document = docx.Document(file_path)
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def extract_text_from_material(file_path: str, content_type: str) -> str:
    if content_type == "application/pdf":
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _extract_xlsx_text(file_path)
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx_text(file_path)
    return Path(file_path).read_text(errors="ignore")
