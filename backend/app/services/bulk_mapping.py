import csv
import io

from openpyxl import load_workbook

# Required column names in the instructor's mapping sheet (case-insensitive,
# whitespace-trimmed). Exact names rather than guessed aliases — easy to
# document in an error message, and avoids silently matching the wrong
# column on a sheet with similarly-named headers.
REQUIRED_COLUMNS = {"student_id_number", "filename"}


def parse_mapping_file(contents: bytes, filename: str) -> list[tuple[str, str]]:
    """Parses a two-column (student_id_number, filename) mapping sheet from
    a .csv or .xlsx file. Raises ValueError with a message safe to show the
    instructor directly."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        header, rows = _read_xlsx(contents)
    elif name.endswith(".csv"):
        header, rows = _read_csv(contents)
    else:
        raise ValueError("Mapping file must be a .csv or .xlsx spreadsheet")
    return _extract_rows(header, rows)


def _read_csv(contents: bytes) -> tuple[list[str], list[dict]]:
    text = contents.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    fieldnames = reader.fieldnames or []
    return list(fieldnames), list(reader)


def _read_xlsx(contents: bytes) -> tuple[list[str], list[dict]]:
    workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        raise ValueError("Mapping spreadsheet is empty")
    fieldnames = [str(h).strip() if h is not None else "" for h in header_row]
    rows = [dict(zip(fieldnames, row)) for row in rows_iter]
    return fieldnames, rows


def _extract_rows(fieldnames: list[str], rows: list[dict]) -> list[tuple[str, str]]:
    normalized_to_actual = {f.strip().lower(): f for f in fieldnames if f}
    missing = REQUIRED_COLUMNS - set(normalized_to_actual)
    if missing:
        raise ValueError(
            "Mapping file must have columns named 'student_id_number' and 'filename' "
            f"(missing: {', '.join(sorted(missing))})"
        )
    id_col = normalized_to_actual["student_id_number"]
    file_col = normalized_to_actual["filename"]

    result = []
    for row in rows:
        student_id_number = str(row.get(id_col) or "").strip()
        filename = str(row.get(file_col) or "").strip()
        if not student_id_number or not filename:
            continue
        result.append((student_id_number, filename))
    return result


def _normalize_filename(filename: str) -> str:
    return filename.strip().lower()


def validate_bulk_mapping(
    mapping_rows: list[tuple[str, str]], uploaded_filenames: list[str]
) -> dict[str, list[str]]:
    """Cross-checks a parsed mapping sheet against the files actually
    uploaded in the same request. Matching is by filename, normalized
    (trimmed, case-insensitive) so a typo in capitalization doesn't block an
    otherwise-correct 160-row batch. Returns a dict of error categories ->
    offending values; empty dict means the batch is safe to apply."""
    errors: dict[str, list[str]] = {}

    seen_ids: dict[str, str] = {}
    duplicate_ids: set[str] = set()
    seen_filenames: dict[str, str] = {}
    duplicate_mapping_filenames: set[str] = set()
    norm_mapping_filenames: set[str] = set()
    for student_id_number, filename in mapping_rows:
        norm_name = _normalize_filename(filename)
        norm_mapping_filenames.add(norm_name)
        if student_id_number in seen_ids and seen_ids[student_id_number] != norm_name:
            duplicate_ids.add(student_id_number)
        seen_ids[student_id_number] = norm_name
        if norm_name in seen_filenames and seen_filenames[norm_name] != student_id_number:
            duplicate_mapping_filenames.add(filename)
        seen_filenames[norm_name] = student_id_number

    if duplicate_ids:
        errors["duplicate_student_ids_in_mapping"] = sorted(duplicate_ids)
    if duplicate_mapping_filenames:
        errors["duplicate_filenames_in_mapping"] = sorted(duplicate_mapping_filenames)

    norm_uploaded = [_normalize_filename(f) for f in uploaded_filenames]
    duplicate_uploads = {
        f for f, n in zip(uploaded_filenames, norm_uploaded) if norm_uploaded.count(n) > 1
    }
    if duplicate_uploads:
        errors["duplicate_uploaded_filenames"] = sorted(duplicate_uploads)

    uploaded_set = set(norm_uploaded)
    unmatched_files = sorted(
        {f for f, n in zip(uploaded_filenames, norm_uploaded) if n not in norm_mapping_filenames}
    )
    unmatched_rows = sorted(
        {filename for _, filename in mapping_rows if _normalize_filename(filename) not in uploaded_set}
    )
    if unmatched_files:
        errors["files_with_no_mapping_row"] = unmatched_files
    if unmatched_rows:
        errors["mapping_rows_with_no_file"] = unmatched_rows

    return errors
